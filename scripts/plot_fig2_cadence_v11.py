"""Fig 2 (two-column layout, 6.75 in wide) — the cadence mechanism (LM, |F|=8, d=128).

Same layout, style and logic as the earlier Fig 2 script; the data come from the
per-run summary and per-epoch curve CSVs written by scripts/analyze-lm-v11.py, so
the figure can be regenerated for any LM run prefix (curve epochs are 1-indexed there).

(a) Final-epoch test ppl at |F|=8: bars = observed means, dots = every seed,
    black line = mean test ppl at the best-validation checkpoint.
(b) Test-ppl training curves at |F|=8, mean over seeds with min-max band.
The panel (a) y-range is 150-310 as before; it is widened (and a note printed)
only if a plotted value would fall outside it.

Usage: python3 scripts/plot_fig2_cadence_v11.py --summary <lm_summary.csv>
           --curves <lm_curves.csv> --out-dir <dir>
Output: <out-dir>/fig2_cadence.{pdf,png[,pptx]} (PNG at 300 dpi).
Every plotted value is printed; lines starting with "VAL|" are full precision.
"""

import argparse
import csv
from collections import defaultdict
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch

ap = argparse.ArgumentParser()
ap.add_argument("--summary", required=True, help="lm_summary.csv from analyze-lm-v11.py")
ap.add_argument("--curves", required=True, help="lm_curves.csv from analyze-lm-v11.py")
ap.add_argument("--out-dir", required=True)
ARGS = ap.parse_args()
OUT = Path(ARGS.out_dir)
OUT.mkdir(parents=True, exist_ok=True)
NAME = "fig2_cadence"

METHODS = ["round_robin", "fixed_cadence_random", "random", "learned_top1"]
C = {"round_robin": "#0072B2", "fixed_cadence_random": "#E69F00",
     "random": "#CC79A7", "learned_top1": "#D55E00"}
LBL = {"round_robin": "Round-\nrobin\nregular\ncanonical\norder",
       "fixed_cadence_random": "Fixed-\ncadence\nregular\nshuffled\norder",
       "random": "Random\nirregular\ni.i.d.",
       "learned_top1": "Learned\ntop-1\nirregular\nlearned"}
CLBL = {"round_robin": "Round-robin (regular)",
        "fixed_cadence_random": "Fixed-cadence random (regular)",
        "random": "Random (irregular)",
        "learned_top1": "Learned top-1 (irregular)"}

plt.rcParams.update({"font.size": 8, "axes.labelsize": 9, "axes.titlesize": 9,
                     "xtick.labelsize": 8, "ytick.labelsize": 8,
                     "legend.fontsize": 8, "pdf.fonttype": 42, "ps.fonttype": 42})


def sd1(v):
    """Sample standard deviation (ddof=1)."""
    return float(np.std(v, ddof=1))


def val(panel, key, v):
    print(f"VAL|fig2|{panel}|{key}|mean={float(np.mean(v))!r}|sd={sd1(v)!r}|n={len(v)}")


summ = defaultdict(list)
bestv = defaultdict(list)
for r in csv.DictReader(open(ARGS.summary)):
    if int(r["backbone"]) == 128 and int(r["F"]) == 8 and r["method"] in METHODS:
        summ[r["method"]].append(float(r["final_test_ppl"]))
        bestv[r["method"]].append(float(r["test_ppl_at_best_val"]))

curves = defaultdict(lambda: defaultdict(list))  # method -> epoch (1-indexed) -> [ppl]
for r in csv.DictReader(open(ARGS.curves)):
    if int(r["backbone"]) == 128 and int(r["F"]) == 8 and r["method"] in METHODS:
        curves[r["method"]][int(r["epoch"])].append(float(r["test_ppl"]))
print(f"LM data: {ARGS.summary}, {ARGS.curves} (backbone d=128, |F|=8)")

fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(6.75, 2.9), layout="constrained",
                               gridspec_kw={"width_ratios": [1.3, 1]})

print("── Fig2(a) encoded table: WikiText-2 |F|=8 test ppl ──")
xs = np.arange(len(METHODS))
BW = 0.62
for i, m in enumerate(METHODS):
    v = summ[m]
    bv = bestv[m]
    print(f"{m}: final-epoch mean={np.mean(v):.1f} sd={sd1(v):.1f} n={len(v)} "
          f"values={[round(x, 1) for x in sorted(v)]}")
    print(f"{m}: best-val-checkpoint mean={np.mean(bv):.1f} sd={sd1(bv):.1f} "
          f"n={len(bv)} values={[round(x, 1) for x in sorted(bv)]}")
    val("a_final", m, v)
    val("a_bestval", m, bv)
    ax1.bar(i, np.mean(v), color=C[m], alpha=0.75, width=BW)
    jit = np.linspace(-0.16, 0.16, len(v))
    ax1.scatter(i + jit, sorted(v), color="black", s=7, zorder=3, alpha=0.75, lw=0)
    ax1.hlines(np.mean(bv), i - BW / 2, i + BW / 2, colors="black", lw=1.6,
               zorder=4)
ax1.set_xticks(xs)
ax1.set_xticklabels([LBL[m] for m in METHODS])
ax1.set_ylabel("Test perplexity (WikiText-2, |F|=8)")
ax1.set_title("(a) Cadence intervention at |F|=8")
YLO, YHI = 150, 310
vals_a = [x for m in METHODS for x in summ[m] + [np.mean(bestv[m])]]
if min(vals_a) < YLO or max(vals_a) > YHI:
    YLO = min(YLO, 10 * np.floor(min(vals_a) / 10) - 10)
    YHI = max(YHI, 10 * np.ceil(max(vals_a) / 10) + 10)
    print(f"NOTE: panel (a) y-range widened to [{YLO}, {YHI}] (data range "
          f"[{min(vals_a):.1f}, {max(vals_a):.1f}])")
ax1.set_ylim(YLO, YHI)
ax1.legend(handles=[Patch(color="grey", alpha=0.5, label="final epoch (mean)"),
                    Line2D([], [], color="black", marker="o", ls="", ms=2.5,
                           alpha=0.75, label="final epoch (per seed)"),
                    Line2D([], [], color="black", lw=1.6,
                           label="best-validation checkpoint (mean)")],
           loc="upper left", frameon=False, handlelength=1.2,
           handletextpad=0.5, borderaxespad=0.3)

print("\n── Fig2(b) encoded table: mean test ppl by epoch (band = seed min–max) ──")
for m in METHODS:
    eps = sorted(curves[m])
    mean = [np.mean(curves[m][e]) for e in eps]
    lo = [np.min(curves[m][e]) for e in eps]
    hi = [np.max(curves[m][e]) for e in eps]
    print(f"{m}: " + " ".join(f"e{e}={mu:.0f}[{a:.0f},{b:.0f}]"
                              for e, mu, a, b in zip(eps, mean, lo, hi)))
    for e in eps:
        print(f"VAL|fig2|b|{m}|epoch={e}|mean={float(np.mean(curves[m][e]))!r}"
              f"|min={float(np.min(curves[m][e]))!r}|max={float(np.max(curves[m][e]))!r}"
              f"|n={len(curves[m][e])}")
    ax2.plot(eps, mean, color=C[m], lw=1.4, marker="o",
             ms=2.5, label=CLBL[m])
    ax2.fill_between(eps, lo, hi, color=C[m], alpha=0.13, lw=0)
ax2.set_xlabel("Epoch")
ax2.set_ylabel("Test perplexity")
ns = sorted({len(curves[m][e]) for m in METHODS for e in curves[m]})
ntxt = f"n={ns[0]}" if len(ns) == 1 else f"n={ns[0]}–{ns[-1]}"
ax2.set_title(f"(b) Test perplexity by epoch, |F|=8\n(band = seed min–max, {ntxt})")
ax2.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=2,
           frameon=False, handlelength=1.5, columnspacing=1.0)
ax2.set_xticks(sorted({e for m in METHODS for e in curves[m]}))

# No suptitle: caption text lives in LaTeX.
for ext in ("pdf", "png"):
    fig.savefig(OUT / f"{NAME}.{ext}", dpi=300)
w, h = fig.get_size_inches()
print(f"\nsaved {OUT}/{NAME}.pdf/.png  size={w:.2f}x{h:.2f} in")

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Fig. 2 — The cadence mechanism (LM, |F|=8)"
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    s.shapes.add_picture(str(OUT / f"{NAME}.png"), Inches(0.4), Inches(1.5),
                         width=Inches(12.533))
    prs.save(OUT / f"{NAME}.pptx")
    print(f"saved {OUT}/{NAME}.pptx")
except ImportError as e:
    print(f"python-pptx not importable ({e}); no PPTX written")
