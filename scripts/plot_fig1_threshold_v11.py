"""Fig 1 (two-column layout, 6.75 in wide) — vision flatness vs LM threshold.

Same layout, style and logic as the earlier Fig 1 script; the LM panel reads its
data from a per-run summary CSV written by scripts/analyze-lm-v11.py, so it can be
regenerated for any LM run prefix.

(a) Vision (S2 natural sweep, CIFAR-100, n=10/cell): mean final test accuracy
    per method vs realized |F| (+-1 sample SD). Data: result/frozen_2026-08-24/s2_summary.csv
    (unchanged data source).
(b) LM (WikiText-2, backbone d=128): mean final-epoch test ppl per method vs |F|
    (+-1 sample SD) with per-seed points at |F|=8. Data: --summary (lm_summary.csv).

Usage: python3 scripts/plot_fig1_threshold_v11.py --summary <lm_summary.csv> --out-dir <dir>
Output: <out-dir>/fig1_threshold.{pdf,png[,pptx]} (PNG at 300 dpi).
Every plotted value is printed; lines starting with "VAL|" are full precision.
"""

import argparse
import csv
from collections import defaultdict
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parent.parent
FROZEN = ROOT / "result" / "frozen_2026-08-24"
ap = argparse.ArgumentParser()
ap.add_argument("--summary", required=True, help="lm_summary.csv from analyze-lm-v11.py")
ap.add_argument("--out-dir", required=True)
ARGS = ap.parse_args()
OUT = Path(ARGS.out_dir)
OUT.mkdir(parents=True, exist_ok=True)
NAME = "fig1_threshold"

# Okabe–Ito colorblind-safe palette (unchanged)
C = {"round_robin": "#0072B2", "gatedmoe_base": "#56B4E9",
     "gatedmoe_g": "#009E73", "learned_top1": "#D55E00",
     "random": "#CC79A7", "fixed_cadence_random": "#E69F00"}
LBL = {"round_robin": "Round-robin (cyclic)", "gatedmoe_base": "Debt-Base (≡ cyclic)",
       "gatedmoe_g": "Debt-Grad", "learned_top1": "Learned top-1",
       "random": "Random (i.i.d.)", "fixed_cadence_random": "Fixed-cadence random"}

plt.rcParams.update({"font.size": 8, "axes.labelsize": 9, "axes.titlesize": 9,
                     "xtick.labelsize": 8, "ytick.labelsize": 8,
                     "legend.fontsize": 8, "pdf.fonttype": 42, "ps.fonttype": 42})


def sd1(v):
    """Sample standard deviation (ddof=1)."""
    return float(np.std(v, ddof=1))


def val(panel, key, v):
    print(f"VAL|fig1|{panel}|{key}|mean={float(np.mean(v))!r}|sd={sd1(v)!r}|n={len(v)}")


# ── (a) vision ───────────────────────────────────────────────────────────────
vis = defaultdict(lambda: defaultdict(list))  # F -> method -> [acc]
for r in csv.DictReader(open(FROZEN / "s2_summary.csv")):
    if r["grid"] == "e1" and r["dataset"] == "c100":
        vis[int(r["expected_F"])][r["method"]].append(float(r["final_test_acc"]))

# ── (b) LM ───────────────────────────────────────────────────────────────────
lm = defaultdict(lambda: defaultdict(list))  # F -> method -> [ppl]
for r in csv.DictReader(open(ARGS.summary)):
    if int(r["backbone"]) == 128:
        lm[int(r["F"])][r["method"]].append(float(r["final_test_ppl"]))
print(f"LM data: {ARGS.summary} (backbone d=128)")

fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(6.75, 2.9), layout="constrained")

vis_methods = ["round_robin", "gatedmoe_base", "gatedmoe_g", "learned_top1", "random"]
Fs = sorted(vis)
print("── Fig1(a) encoded table: CIFAR-100 final test acc, mean ± sample SD (n) ──")
for F in Fs:
    print(f"|F|={F}: " + " ".join(
        f"{m}={np.mean(vis[F][m]):.4f}±{sd1(vis[F][m]):.4f}(n{len(vis[F][m])})"
        for m in vis_methods))
    for m in vis_methods:
        val("a", f"F={F}|{m}", vis[F][m])

# Floor = mean over seeds of the per-seed max-min spread across methods at
# |F|=0 (the matched no-routing spread). Shown for scale only.
per_seed = defaultdict(dict)
for r in csv.DictReader(open(FROZEN / "s2_summary.csv")):
    if r["grid"] == "e1" and r["dataset"] == "c100" and int(r["expected_F"]) == 0:
        per_seed[int(r["seed"])][r["method"]] = float(r["final_test_acc"])
spreads0 = [max(a.values()) - min(a.values())
            for a in per_seed.values() if len(a) == len(vis_methods)]
floor_spread = float(np.mean(spreads0))
ymid = float(np.mean([np.mean(vis[0][m]) for m in vis_methods]))
print(f"floor: mean per-seed matched spread at |F|=0 = {floor_spread:.4f} "
      f"({floor_spread * 100:.2f} pp) over {len(spreads0)} seeds; "
      f"band centre = {ymid:.4f}; band = [{ymid - floor_spread / 2:.4f}, "
      f"{ymid + floor_spread / 2:.4f}]")
print(f"VAL|fig1|a|band|height={floor_spread!r}|centre={ymid!r}|nseeds={len(spreads0)}")
for m in vis_methods:
    ys = [np.mean(vis[F][m]) for F in Fs]
    sds = [sd1(vis[F][m]) for F in Fs]
    ax1.errorbar(Fs, ys, yerr=sds, color=C[m], label=LBL[m], marker="o",
                 ms=3, lw=1.2, elinewidth=0.8, capsize=1.5, alpha=0.9)
ax1.axhspan(ymid - floor_spread / 2, ymid + floor_spread / 2, color="grey",
            alpha=0.20, lw=0,
            label=f"no-routing spread at |F|=0 ({floor_spread * 100:.2f} pp), for scale")
ax1.set_xlabel("Feasible-set size |F| (budget-determined)")
ax1.set_ylabel("Final test accuracy (CIFAR-100)")
ax1.set_title("(a) Vision: no resolved router\ndifference at any |F|")
ax1.set_xticks(Fs)

lm_methods = ["round_robin", "gatedmoe_base", "gatedmoe_g", "learned_top1", "random"]
lFs = sorted(lm)
print("\n── Fig1(b) encoded table: WikiText-2 final test ppl, mean ± sample SD (n) ──")
for F in lFs:
    print(f"|F|={F}: " + " ".join(
        f"{m}={np.mean(lm[F][m]):.1f}±{sd1(lm[F][m]):.1f}(n{len(lm[F][m])})"
        for m in lm_methods if lm[F].get(m)))
    for m in lm_methods:
        if lm[F].get(m):
            val("b", f"F={F}|{m}", lm[F][m])
for m in lm_methods:
    xs = [F for F in lFs if lm[F].get(m)]
    ys = [np.mean(lm[F][m]) for F in xs]
    sds = [sd1(lm[F][m]) for F in xs]
    ax2.errorbar(xs, ys, yerr=sds, color=C[m], label=LBL[m], marker="o",
                 ms=3, lw=1.2, elinewidth=0.8, capsize=1.5, alpha=0.9)
    # per-seed points at |F|=8 (jittered on x for visibility, values exact)
    if lm[8].get(m):
        jit = np.linspace(-0.18, 0.18, len(lm[8][m]))
        ax2.scatter(8 + jit, lm[8][m], color=C[m], s=5, alpha=0.5, zorder=3, lw=0)
        print(f"per-seed |F|=8 {m}: {[round(x, 1) for x in lm[8][m]]}")
ax2.set_xlabel("Feasible-set size |F| (budget-determined)")
ax2.set_ylabel("Final test perplexity (WikiText-2)")
ax2.set_title("(b) LM, final epoch: router gaps\nopen at high |F|")
ax2.set_xticks(lFs)

# One shared legend below both panels (no data occlusion), original order.
handles, labels = ax1.get_legend_handles_labels()
print("legend order:", labels)
fig.legend(handles, labels, loc="outside lower center", ncol=3, frameon=False,
           columnspacing=1.2, handlelength=1.8)
# No suptitle: caption text lives in LaTeX.
for ext in ("pdf", "png"):
    fig.savefig(OUT / f"{NAME}.{ext}", dpi=300)
w, h = fig.get_size_inches()
print(f"\nsaved {OUT}/{NAME}.pdf/.png  size={w:.2f}x{h:.2f} in")

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Fig. 1 — When does router choice matter?"
    s.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    s.shapes.add_picture(str(OUT / f"{NAME}.png"), Inches(0.4), Inches(1.5),
                         width=Inches(12.533))
    prs.save(OUT / f"{NAME}.pptx")
    print(f"saved {OUT}/{NAME}.pptx")
except ImportError as e:
    print(f"python-pptx not importable ({e}); no PPTX written")
